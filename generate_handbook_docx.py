# -*- coding: utf-8 -*-
"""
Generate docx for Deep AI Theory & Biomedical Engineering Handbook
Output: docs/Ly_thuyet_Chuyen_sau_va_Kien_thuc_AI_NEURO_ECG.docx
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, title, text, bg_hex="F0F4F8", border_hex="0A2540"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="{border_hex}"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    run_t = p.add_run(f"★ {title}\n")
    run_t.font.name = 'Calibri'
    run_t.font.size = Pt(11)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(10, 37, 64)
    
    run_c = p.add_run(text)
    run_c.font.name = 'Calibri'
    run_c.font.size = Pt(10)
    run_c.font.color.rgb = RGBColor(40, 45, 55)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(4)

def build_handbook():
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)
        
    NAVY = RGBColor(10, 37, 64)
    BLUE = RGBColor(0, 114, 206)
    DARK = RGBColor(33, 37, 41)
    MUTED = RGBColor(100, 110, 120)
    
    # Title
    p_pre = doc.add_paragraph()
    r_pre = p_pre.add_run("SAMSUNG INNOVATION CAMPUS (SIC) 2026 — ĐẠI HỌC BÁCH KHOA HÀ NỘI")
    r_pre.font.name = 'Calibri'
    r_pre.font.size = Pt(10)
    r_pre.font.bold = True
    r_pre.font.color.rgb = BLUE
    
    p_title = doc.add_paragraph()
    r_title = p_title.add_run("BÁCH KHOA TOÀN THƯ LÝ THUYẾT VÀ KIẾN THỨC CHUYÊN SÂU AI & Y SINH\nHỆ THỐNG GIÁM SÁT ECG THỜI GIAN THỰC NEURO-ECG")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(17)
    r_title.font.bold = True
    r_title.font.color.rgb = NAVY
    
    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run("Cẩm nang bảo vệ đồ án tốt nghiệp dành riêng cho Báo cáo viên chính: Nguyễn Đào Nam Hải (Team Leader)")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = MUTED
    
    # Read Markdown content and build clean paragraphs
    with open('docs/Ly_thuyet_Chuyen_sau_va_Kien_thuc_AI_NEURO_ECG.md', 'r', encoding='utf-8') as f:
        md_text = f.read()
        
    sections = md_text.split('\n## ')
    # Header part
    p_intro = doc.add_paragraph()
    p_intro.add_run("Tài liệu này tổng hợp toàn bộ nền tảng toán học, giải phẫu sinh lý học tim mạch, cơ chế các mô hình học sâu 1D, kỹ thuật chống Data Leakage, bản đồ giải thích 1D Grad-CAM, lượng tử hóa ONNX INT8, và các câu hỏi phản biện chuyên sâu từ Hội đồng SIC.").font.size = Pt(10.5)
    
    for sec in sections[1:]:
        lines = sec.strip().split('\n')
        h2_title = lines[0].strip()
        body = '\n'.join(lines[1:]).strip()
        
        p_h2 = doc.add_paragraph()
        p_h2.paragraph_format.space_before = Pt(14)
        p_h2.paragraph_format.space_after = Pt(4)
        r_h2 = p_h2.add_run(h2_title.upper())
        r_h2.font.name = 'Calibri'
        r_h2.font.size = Pt(13)
        r_h2.font.bold = True
        r_h2.font.color.rgb = NAVY
        
        # Parse callouts or text
        paragraphs_in_sec = body.split('\n\n')
        for par in paragraphs_in_sec:
            par = par.strip()
            if not par:
                continue
            if par.startswith('### '):
                p_h3 = doc.add_paragraph()
                p_h3.paragraph_format.space_before = Pt(8)
                p_h3.paragraph_format.space_after = Pt(2)
                r_h3 = p_h3.add_run(par[4:].strip())
                r_h3.font.name = 'Calibri'
                r_h3.font.size = Pt(11.5)
                r_h3.font.bold = True
                r_h3.font.color.rgb = BLUE
            elif par.startswith('> '):
                # Quote box
                add_callout(doc, "LUẬN ĐIỂM TRẢ LỜI PHẢN BIỆN CHUẨN", par.replace('> ', '').strip(), bg_hex="FFFDF5", border_hex="D97706")
            elif par.startswith('|'):
                # Markdown table -> simple text block
                p_tbl = doc.add_paragraph()
                p_tbl.paragraph_format.line_spacing = 1.15
                r_tbl = p_tbl.add_run(par)
                r_tbl.font.name = 'Consolas'
                r_tbl.font.size = Pt(8.5)
                r_tbl.font.color.rgb = DARK
            else:
                p_p = doc.add_paragraph()
                p_p.paragraph_format.line_spacing = 1.15
                p_p.paragraph_format.space_after = Pt(4)
                r_p = p_p.add_run(par)
                r_p.font.name = 'Calibri'
                r_p.font.size = Pt(10)
                r_p.font.color.rgb = DARK
                
    out_docx = r"docs/Ly_thuyet_Chuyen_sau_va_Kien_thuc_AI_NEURO_ECG.docx"
    doc.save(out_docx)
    print(f"Successfully generated handbook at {out_docx}")

if __name__ == '__main__':
    build_handbook()
